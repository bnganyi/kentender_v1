#!/usr/bin/env python3
"""Bring the Documentation Control workbook into line with the baseline register.

The register (`docs/mvp-1-r1/98_work_progress/KenTender_Baseline_Register.yaml`, JSON) is the canonical source; the workbook
(`KenTender_Documentation_Control_Register.xlsx`) is its companion view. This script rewrites every table row from the register,
adds rows the workbook lacks, extends each table and the Overview formulas to the new row counts, keeps the workbook's own styling,
and then reads the workbook back and compares it with the register cell by cell. It exits 1 if any cell still differs.

    /home/midasuser/frappe-bench/env/bin/python scripts/sync_control_workbook.py [--check]

`--check` writes nothing and only reports differences. Needs `openpyxl` (the bench virtualenv has it).
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

PACK = Path(__file__).resolve().parents[1] / "docs" / "mvp-1-r1" / "98_work_progress"
REGISTER = PACK / "KenTender_Baseline_Register.yaml"
WORKBOOK = PACK / "KenTender_Documentation_Control_Register.xlsx"
HEADER_ROW, FIRST_ROW = 5, 6


def source_of(doc: dict) -> str:
	name = doc.get("filename") or ""
	if not name:
		return ""
	return f"Repository: {name}" if "/" in name else name


def approval_date(value: str):
	return datetime.datetime.strptime(value, "%Y-%m-%d") if value else None


def history(doc: dict, key: str) -> str:
	"""A proposal still awaiting the Owner: once its proposal_status reads Approved, or it is the document's own version, it is not a proposal."""
	record = doc.get("approval_history") or {}
	if str(record.get("proposal_status") or "").startswith("Approved") or str(record.get("proposed_version") or "") == str(doc.get("version") or ""):
		return ""
	return record.get(key) or ""


# sheet -> (register key, id field, {column header: callable(register item) -> value})
SHEETS = {
	"Documents": ("documents", "document_id", {
		"Document ID": lambda d: d["document_id"], "Title": lambda d: d["title"], "Version": lambda d: d["version"], "Module": lambda d: d["module"],
		"Category": lambda d: d["category"], "Requirement status": lambda d: d["requirement_status"], "Approval date": lambda d: approval_date(d.get("approval_date") or ""),
		"Supersedes / basis": lambda d: d.get("supersedes") or "", "Implementation": lambda d: d.get("implementation_status") or "", "Verification": lambda d: d.get("verification_status") or "",
		"Release": lambda d: d.get("release_status") or "", "Normative": lambda d: d.get("normative") or "", "Scope": lambda d: d.get("scope") or "", "Dependencies": lambda d: d.get("dependencies") or "",
		"Required action": lambda d: d.get("action") or "", "Source": source_of, "Proposed version": lambda d: history(d, "proposed_version"), "Proposed file": lambda d: history(d, "proposed_filename"),
	}),
	"Interfaces": ("interfaces", "interface_id", {
		"Interface ID": lambda d: d["interface_id"], "Upstream owner": lambda d: d["upstream"], "Downstream consumer": lambda d: d["downstream"], "Contract / artifact": lambda d: d["contract"],
		"Requirement status": lambda d: d["status"], "Implementation": lambda d: d.get("implementation_status") or "", "Test status": lambda d: d.get("test_status") or "",
		"Owning documents": lambda d: d.get("owner_doc") or "", "Failure rule": lambda d: d.get("failure_rule") or "",
	}),
	"Decisions": ("decisions", "decision_id", {
		"Decision ID": lambda d: d["decision_id"], "Decision": lambda d: d["decision"], "Status": lambda d: d["status"], "Owning document": lambda d: d.get("owner_doc") or "",
		"Implementation effect": lambda d: d.get("implementation_effect") or "", "Design effect": lambda d: d.get("design_effect") or "",
	}),
	"Delivery": ("delivery_items", "item_id", {
		"Item ID": lambda d: d["item_id"], "Area": lambda d: d["area"], "Required change": lambda d: d["change"], "Functional outcome": lambda d: d.get("functional") or "",
		"Design outcome": lambda d: d.get("design") or "", "Owning document": lambda d: d.get("owner_doc") or "", "Requirement status": lambda d: d["requirement_status"],
		"Release gate": lambda d: d.get("release_gate") or "", "Implementation": lambda d: d.get("implementation_status") or "", "Verification": lambda d: d.get("verification_status") or "",
		"Evidence / notes": lambda d: d.get("evidence") or "",
	}),
	"Future Work": ("future_work", "future_id", {
		"Future ID": lambda d: d["future_id"], "Capability": lambda d: d["capability"], "Planned wave": lambda d: d.get("planned_wave") or "", "Dependency": lambda d: d.get("dependency") or "",
		"Status": lambda d: d["status"], "Current boundary": lambda d: d.get("boundary") or "", "Trigger to start": lambda d: d.get("trigger") or "",
	}),
	"Audit": ("audit_findings", "finding_id", {
		"Finding ID": lambda d: d["finding_id"], "Severity": lambda d: d.get("severity") or "", "Type": lambda d: d.get("type") or "", "Finding": lambda d: d.get("finding") or "",
		"Status": lambda d: d.get("status") or "", "Owner": lambda d: d.get("owner") or "", "Required action": lambda d: d.get("action") or "", "Blocks": lambda d: d.get("blocks") or "",
		"Evidence": lambda d: d.get("evidence") or "",
	}),
}


def norm(value) -> str:
	if value is None:
		return ""
	if isinstance(value, datetime.datetime):
		return value.strftime("%Y-%m-%d")
	if isinstance(value, int) and 40000 < value < 60000:  # a date stored as an Excel serial number (a few cells were typed that way)
		return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=value)).strftime("%Y-%m-%d")
	return re.sub(r"\s+", " ", str(value)).strip()


def headers(ws) -> list[tuple[int, str]]:
	return [(c.column, c.value) for c in ws[HEADER_ROW] if c.value]


def sync_sheet(ws, key: str, idf: str, columns: dict, register: dict, write: bool) -> tuple[int, int, int]:
	cols = headers(ws)
	missing = [h for _c, h in cols if h not in columns]
	if missing:
		raise SystemExit(f"{ws.title}: no register mapping for column(s) {missing}")
	items = register[key]
	existing = {ws.cell(r, 1).value: r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value}
	last = max(existing.values()) if existing else FIRST_ROW - 1
	changed = added = 0
	for item in items:
		ident = item[idf]
		row = existing.get(ident)
		if row is None:
			last += 1
			row, added = last, added + 1
			if write:
				for col, _h in cols:
					ws.cell(row, col)._style = ws.cell(last - 1, col)._style
				if ws.row_dimensions[last - 1].height:
					ws.row_dimensions[row].height = ws.row_dimensions[last - 1].height
			existing[ident] = row
		for col, header in cols:
			wanted = columns[header](item)
			cell = ws.cell(row, col)
			if norm(cell.value) != norm(wanted):
				changed += 1
				if write:
					cell.value = wanted if wanted != "" else None
	if write and ws.tables:
		table = next(iter(ws.tables.values()))
		ref = f"A{HEADER_ROW}:{openpyxl.utils.get_column_letter(cols[-1][0])}{last}"
		table.ref = ref
		if table.autoFilter is not None:
			table.autoFilter.ref = ref
	return changed, added, last


def verify(wb, register: dict) -> list[str]:
	problems = []
	for name, (key, idf, columns) in SHEETS.items():
		ws = wb[name]
		by_id = {ws.cell(r, 1).value: r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value}
		for item in register[key]:
			row = by_id.get(item[idf])
			if row is None:
				problems.append(f"{name}: {item[idf]} is missing")
				continue
			for col, header in headers(ws):
				if norm(ws.cell(row, col).value) != norm(columns[header](item)):
					problems.append(f"{name}: {item[idf]} / {header} differs")
		extra = set(by_id) - {i[idf] for i in register[key]}
		problems += [f"{name}: {e} is not in the register" for e in sorted(extra)]
	return problems


def main() -> int:
	parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
	parser.add_argument("--check", action="store_true", help="report differences; write nothing")
	args = parser.parse_args()
	register = json.loads(REGISTER.read_text(encoding="utf-8"))
	wb = openpyxl.load_workbook(WORKBOOK)
	lasts = {}
	for name, (key, idf, columns) in SHEETS.items():
		changed, added, last = sync_sheet(wb[name], key, idf, columns, register, write=not args.check)
		lasts[name] = last
		print(f"{name}: {changed} cell(s) {'differ' if args.check else 'updated'}, {added} row(s) {'missing' if args.check else 'added'}")
	if not args.check:
		approved = max((d["approval_date"] for d in register["documents"] if d.get("approval_date")), default="")
		note = f"Approved version set: {approved}. Proposals up to {register['as_of']} are listed beside it. Design work remains open; CM and template walkthrough deferred."
		for name in SHEETS:
			if str(wb[name]["A2"].value or "").startswith("Approved version set"):
				wb[name]["A2"].value = note
		overview = wb["Overview"]
		d, dl = lasts["Documents"], lasts["Delivery"]
		overview["A6"].value = f"=COUNTA(Documents!A6:A{d})"
		overview["C6"].value = f'=COUNTIF(Documents!F6:F{d},"Approved requirement")'
		overview["E6"].value = f'=COUNTA(Documents!Q6:Q{d})+COUNTIF(Documents!F6:F{d},"Project Owner review")'
		overview["G6"].value = f'=COUNTIF(Delivery!I6:I{dl},"Not started")+COUNTIF(Delivery!I6:I{dl},"Planned")+COUNTIF(Delivery!I6:I{dl},"Pending")'
		overview["A2"].value = f"Register as at {register['as_of']}: approved requirements through {approved}; later proposals are listed in Documents."
		wb.save(WORKBOOK)
		wb = openpyxl.load_workbook(WORKBOOK)
	problems = verify(wb, register)
	for line in problems[:40]:
		print("  " + line)
	print("workbook matches the register" if not problems else f"{len(problems)} difference(s) remain")
	return 1 if problems else 0


if __name__ == "__main__":
	sys.exit(main())
